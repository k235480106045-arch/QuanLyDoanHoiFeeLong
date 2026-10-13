from flask import Flask, render_template, request, redirect, url_for, session, send_file, Response
import mysql.connector
import os
import io
import csv
from werkzeug.utils import secure_filename
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

app = Flask(__name__)
app.secret_key = 'doanhoi_tnut_secret_key'
# Cấu hình thư mục upload
UPLOAD_FOLDER = 'static/uploads'
IMAGE_FOLDER = 'static/images/uploads'
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(IMAGE_FOLDER, exist_ok=True)
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER
app.config['IMAGE_UPLOAD_FOLDER'] = IMAGE_FOLDER

AVATAR_FOLDER = 'static/avatars'
os.makedirs(AVATAR_FOLDER, exist_ok=True)
app.config['AVATAR_FOLDER'] = AVATAR_FOLDER

# ================= HÀM KIỂM TRA PHÂN QUYỀN =================
def require_login():
    if 'user_id' not in session:
        return redirect(url_for('index'))
    return None

def require_role(roles):
    def decorator(f):
        def wrapper(*args, **kwargs):
            if 'user_id' not in session:
                return redirect(url_for('index'))
            if session.get('role') not in roles:
                return f"<script>alert('Bạn không có quyền thực hiện thao tác này!'); window.location.href='/home';</script>"
            return f(*args, **kwargs)
        wrapper.__name__ = f.__name__
        return wrapper
    return decorator

# ================= HÀM KẾT NỐI CSDL =================
def get_db_connection():
    return mysql.connector.connect(
        host='127.0.0.1',
        user='root',
        password='',
        database='quanlydtn_hsv'
    )

# ================= TÀI KHOẢN & ĐĂNG NHẬP =================

@app.route('/')
def index():
    if 'user_id' in session:
        return redirect(url_for('home'))
    return render_template('login.html')

@app.route('/login', methods=['POST'])
def login():
    username = request.form['username'].strip()
    password = request.form['password'].strip()
    ip = request.remote_addr or ''
    ua = (request.user_agent.string or '')[:500]
    
    def ghi_log(tt, vt='', hoten='', ct=''):
        try:
            lc = get_db_connection(); lc.cursor().execute(
                "INSERT INTO login_log (username, vai_tro, ho_ten, ip_address, user_agent, trang_thai, chi_tiet) VALUES (%s,%s,%s,%s,%s,%s,%s)",
                (username, vt, hoten, ip, ua, tt, ct)
            ); lc.commit(); lc.close()
        except: pass
    
    try:
        conn = get_db_connection()
        cursor = conn.cursor(dictionary=True)
        
        query = "SELECT * FROM nguoi_dung WHERE TenDangNhap = %s AND MatKhau = %s"
        cursor.execute(query, (username, password))
        user = cursor.fetchone()
        
        cursor.close()
        conn.close()
        
        if user:
            import datetime
            ho_ten_user = user.get('ho_ten') or user.get('TenDangNhap', '')
            ghi_log('thanh_cong', user['VaiTro'], ho_ten_user, f"Đăng nhập thành công")
            # Cập nhật lần đăng nhập cuối
            try:
                uc = get_db_connection()
                uc.cursor().execute("UPDATE nguoi_dung SET last_login = %s WHERE MaNguoiDung = %s", (datetime.datetime.now(), user['MaNguoiDung']))
                uc.commit(); uc.close()
            except: pass
            
            session['user_id'] = user['MaNguoiDung']
            session['username'] = user['TenDangNhap']
            session['role'] = user['VaiTro']
            session['avatar'] = user.get('avatar') or ''
            session['masv'] = user.get('MaSV') or ''
            
            if user['LaLanDangNhapDau']:
                return redirect(url_for('change_password'))
            else:
                return redirect(url_for('home')) 
        else:
            ghi_log('that_bai', ct='Sai tài khoản hoặc mật khẩu')
            return render_template('login.html', error="Tài khoản hoặc mật khẩu không chính xác!")
            
    except Exception as e:
        return f"Lỗi kết nối CSDL: {e}"

@app.route('/change-password', methods=['GET', 'POST'])
def change_password():
    if 'user_id' not in session:
        return redirect(url_for('index'))
        
    if request.method == 'POST':
        new_password = request.form['new_password']
        confirm_password = request.form['confirm_password']
        
        if new_password != confirm_password:
            return render_template('change_password.html', error="Mật khẩu nhập lại không khớp!")
        
        try:
            user_id = session['user_id']
            conn = get_db_connection()
            cursor = conn.cursor()
            
            query = "UPDATE nguoi_dung SET MatKhau = %s, LaLanDangNhapDau = 0 WHERE MaNguoiDung = %s"
            cursor.execute(query, (new_password, user_id))
            conn.commit()
            
            cursor.close()
            conn.close()
            
            session.clear() 
            return render_template('login.html', error="Đổi mật khẩu thành công! Vui lòng đăng nhập lại.")
            
        except Exception as e:
            return f"Lỗi: {e}"
            
    return render_template('change_password.html')

@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('index'))

# ================= TRANG CHỦ & BÀI VIẾT =================

@app.route('/home')
def home():
    if 'user_id' not in session: return redirect(url_for('index'))
    
    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)
    cursor.execute("SELECT * FROM bai_viet ORDER BY ngay_tao DESC")
    bai_viet_list = cursor.fetchall()
    cursor.close(); conn.close()
    
    intro_content = get_intro()
    
    return render_template('home.html', bai_viet_list=bai_viet_list, intro_content=intro_content)

@app.route('/tao-bai-viet', methods=['POST'])
def tao_bai_viet():
    if 'user_id' not in session or session.get('role') not in ['admin', 'bch']:
        return f"<script>alert('Bạn không có quyền thực hiện thao tác này!'); window.location.href='/home';</script>"
        
    noi_dung = request.form.get('noi_dung', '')
    hinh_anh = request.files.get('hinh_anh') 
    filename = None
    
    if hinh_anh and hinh_anh.filename != '':
        filename = secure_filename(hinh_anh.filename)
        hinh_anh.save(os.path.join(app.config['IMAGE_UPLOAD_FOLDER'], filename))
        
    # LƯU VÀO DATABASE
    try:
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute("INSERT INTO bai_viet (noi_dung, hinh_anh) VALUES (%s, %s)", (noi_dung, filename))
        conn.commit()
        cursor.close(); conn.close()
    except Exception as e:
        print(f"Lỗi lưu bài: {e}")
    return redirect(url_for('home'))


@app.route('/sua-bai-viet-trang-chu/<int:ma_bv>', methods=['POST'])
def sua_bai_viet_trang_chu(ma_bv):
    if 'user_id' not in session or session.get('role') not in ['admin', 'bch']:
        return f"<script>alert('Bạn không có quyền!'); window.location.href='/home';</script>"
    noi_dung = request.form.get('noi_dung', '')
    try:
        conn = get_db_connection(); cursor = conn.cursor()
        cursor.execute("UPDATE bai_viet SET noi_dung = %s WHERE id = %s", (noi_dung, ma_bv))
        hinh_anh = request.files.get('hinh_anh')
        if hinh_anh and hinh_anh.filename != '':
            filename = secure_filename(hinh_anh.filename)
            hinh_anh.save(os.path.join(app.config['IMAGE_UPLOAD_FOLDER'], filename))
            cursor.execute("UPDATE bai_viet SET hinh_anh = %s WHERE id = %s", (filename, ma_bv))
        conn.commit(); cursor.close(); conn.close()
        return "<script>alert('Cập nhật bài viết thành công!'); window.location.href='/home';</script>"
    except Exception as e:
        return f"<script>alert('Lỗi: {e}'); window.location.href='/home';</script>"


@app.route('/xoa-bai-viet-trang-chu/<int:ma_bv>')
def xoa_bv_trang_chu(ma_bv):
    if 'user_id' not in session or session.get('role') not in ['admin', 'bch']:
        return f"<script>alert('Bạn không có quyền!'); window.location.href='/home';</script>"
    try:
        conn = get_db_connection(); cursor = conn.cursor()
        cursor.execute("DELETE FROM bai_viet WHERE id = %s", (ma_bv,))
        conn.commit(); cursor.close(); conn.close()
        return "<script>alert('Đã xóa bài viết!'); window.location.href='/home';</script>"
    except Exception as e:
        return f"<script>alert('Lỗi: {e}'); window.location.href='/home';</script>"


@app.route('/xoa-anh-bai-viet-trang-chu/<int:ma_bv>')
def xoa_anh_bv_trang_chu(ma_bv):
    if 'user_id' not in session or session.get('role') not in ['admin', 'bch']:
        return f"<script>alert('Bạn không có quyền!'); window.location.href='/home';</script>"
    try:
        conn = get_db_connection(); cursor = conn.cursor()
        cursor.execute("UPDATE bai_viet SET hinh_anh = NULL WHERE id = %s", (ma_bv,))
        conn.commit(); cursor.close(); conn.close()
        return "<script>alert('Đã xóa ảnh!'); window.location.href='/home';</script>"
    except Exception as e:
        return f"<script>alert('Lỗi: {e}'); window.location.href='/home';</script>"


@app.route('/delete-avatar')
def delete_avatar():
    if 'user_id' not in session: return redirect(url_for('index'))
    try:
        conn = get_db_connection(); cursor = conn.cursor()
        cursor.execute("UPDATE nguoi_dung SET avatar = NULL WHERE MaNguoiDung = %s", (session['user_id'],))
        conn.commit(); cursor.close(); conn.close()
        session['avatar'] = ''
        return "<script>alert('Đã xóa ảnh đại diện!'); window.location.href='/profile';</script>"
    except Exception as e:
        return f"<script>alert('Lỗi: {e}'); window.location.href='/profile';</script>"

@app.route('/doi-mat-khau', methods=['POST'])
def doi_mat_khau():
    if 'user_id' not in session: return redirect(url_for('index'))
    old = request.form.get('old_password')
    new = request.form.get('new_password')
    confirm = request.form.get('confirm_password')
    if new != confirm:
        return "<script>alert('Mật khẩu mới không khớp!'); window.location.href='/profile';</script>"
    try:
        conn = get_db_connection(); cursor = conn.cursor(dictionary=True)
        cursor.execute("SELECT MatKhau FROM nguoi_dung WHERE MaNguoiDung = %s", (session['user_id'],))
        u = cursor.fetchone()
        if u and u['MatKhau'] != old:
            return "<script>alert('Mật khẩu cũ không đúng!'); window.location.href='/profile';</script>"
        cursor.execute("UPDATE nguoi_dung SET MatKhau = %s WHERE MaNguoiDung = %s", (new, session['user_id']))
        conn.commit(); cursor.close(); conn.close()
        return "<script>alert('Đổi mật khẩu thành công!'); window.location.href='/profile';</script>"
    except Exception as e:
        return f"<script>alert('Lỗi: {e}'); window.location.href='/profile';</script>"

# ================= TRANG CÁ NHÂN & AVATAR =================

@app.route('/profile')
def profile():
    r = require_login()
    if r: return r

    conn = get_db_connection()
    cursor = conn.cursor(dictionary=True)
    cursor.execute("SELECT * FROM nguoi_dung WHERE MaNguoiDung = %s", (session['user_id'],))
    user = cursor.fetchone()

    # Lấy thông tin sinh viên nếu có
    sv_info = None
    if user.get('MaSV'):
        cursor.execute("SELECT * FROM sinh_vien WHERE MaSV = %s", (user['MaSV'],))
        sv_info = cursor.fetchone()

    # Thống kê
    ma_sv = user.get('MaSV') or ''
    cursor.execute("SELECT COUNT(*) as cnt FROM sv_tham_gia_hd WHERE MaSV = %s", (ma_sv,))
    da_tham_gia = cursor.fetchone()['cnt']

    cursor.execute("SELECT COUNT(*) as cnt FROM sv_khen_thuong WHERE MaSV = %s", (ma_sv,))
    so_khen_thuong = cursor.fetchone()['cnt']

    cursor.execute("SELECT COALESCE(SUM(Diem), 0) as tong FROM diem_ren_luyen WHERE MaSV = %s", (ma_sv,))
    tong_drl = cursor.fetchone()['tong']

    # Chi tiết hoạt động đã tham gia
    cursor.execute("""
        SELECT hd.TenHD, hd.NgayBatDau, hd.NgayKetThuc, hd.DiemCong, hd.HocKy, hd.NamHoc,
               stg.VaiTro, lhd.TenLoaiHD
        FROM sv_tham_gia_hd stg
        JOIN hoat_dong hd ON stg.MaHD = hd.MaHD
        LEFT JOIN loai_hoat_dong lhd ON hd.MaLoaiHD = lhd.MaLoaiHD
        WHERE stg.MaSV = %s ORDER BY hd.NgayBatDau DESC
    """, (ma_sv,))
    ds_hoat_dong = cursor.fetchall()

    # Chi tiết điểm rèn luyện theo học kỳ
    cursor.execute("""
        SELECT HocKy, NamHoc, Diem FROM diem_ren_luyen
        WHERE MaSV = %s ORDER BY NamHoc DESC, HocKy DESC
    """, (ma_sv,))
    ds_diem_ren_luyen = cursor.fetchall()

    # Chi tiết khen thưởng
    cursor.execute("""
        SELECT kt.NgayKhenThuong, kt.LyDo, kt.TrangThai, kt.FileMinhChung,
               dm.TenDanhMuc
        FROM sv_khen_thuong kt
        LEFT JOIN danh_muc_khen_thuong dm ON kt.MaDanhMucKT = dm.MaDanhMucKT
        WHERE kt.MaSV = %s ORDER BY kt.NgayKhenThuong DESC
    """, (ma_sv,))
    ds_khen_thuong = cursor.fetchall()

    thong_ke = {
        'so_luong_hoat_dong': da_tham_gia,
        'so_luong_khen_thuong': so_khen_thuong,
        'tong_diem_ren_luyen': tong_drl,
        'ds_hoat_dong': ds_hoat_dong,
        'ds_diem_ren_luyen': ds_diem_ren_luyen,
        'ds_khen_thuong': ds_khen_thuong
    }

    # Xây dựng user_info dict từ user + sv_info
    user_info = {
        'HoTen': user.get('ho_ten') or (f"{sv_info['HoLotSV']} {sv_info['TenSV']}" if sv_info else ''),
        'MaSo': user.get('MaSV') or user.get('TenDangNhap', ''),
        'Email': user.get('email_personal') or (sv_info['Email'] if sv_info else ''),
        'TenDangNhap': user.get('TenDangNhap', session.get('username', '')),
        'GioiTinh': sv_info.get('GioiTinh', '') if sv_info else '',
        'TenLop': sv_info.get('TenLop', '') if sv_info else '',
        'NgaySinh': str(sv_info.get('NgaySinh', '')) if sv_info else '',
        'SoDienThoai': sv_info.get('SDT', '') if sv_info else '',
        'DiaChi': sv_info.get('QueQuan', '') if sv_info else '',
        'ChucVu': '',
        'bio': user.get('bio', '')
    }

    avatar_url = f"/static/avatars/{user['avatar']}" if user.get('avatar') else ''

    cursor.close(); conn.close()
    return render_template('profile.html', user_info=user_info, thong_ke=thong_ke,
                           avatar_url=avatar_url)


@app.route('/upload-avatar', methods=['POST'])
def upload_avatar():
    if 'user_id' not in session: return redirect(url_for('index'))
    file = request.files.get('avatar')
    try:
        if file and file.filename:
            original_name = secure_filename(file.filename)
            ext = original_name.rsplit('.', 1)[1].lower() if '.' in original_name else 'png'
            safe_username = secure_filename(session.get('username', 'user'))
            filename = f"user_{session['user_id']}_{safe_username}.{ext}"
            save_path = os.path.join(app.config['AVATAR_FOLDER'], filename)
            file.save(save_path)
            conn = get_db_connection()
            cursor = conn.cursor()
            cursor.execute("UPDATE nguoi_dung SET avatar = %s WHERE MaNguoiDung = %s", (filename, session['user_id']))
            conn.commit()
            cursor.close(); conn.close()
            session['avatar'] = filename
    except Exception as e:
        return f"<script>alert('Lỗi upload ảnh: {e}'); window.location.href='/profile';</script>"
    return redirect(url_for('profile'))


@app.route('/cap-nhat-profile', methods=['POST'])
def cap_nhat_profile():
    if 'user_id' not in session: return redirect(url_for('index'))
    ho_ten = request.form.get('ho_ten')
    email = request.form.get('email_personal')
    bio = request.form.get('bio')
    try:
        conn = get_db_connection(); cursor = conn.cursor()
        cursor.execute("UPDATE nguoi_dung SET ho_ten=%s, email_personal=%s, bio=%s WHERE MaNguoiDung=%s",
                       (ho_ten, email, bio, session['user_id']))
        conn.commit(); cursor.close(); conn.close()
        return "<script>alert('Cập nhật thành công!'); window.location.href='/profile';</script>"
    except Exception as e:
        return f"<script>alert('Lỗi: {e}'); window.location.href='/profile';</script>"


# ================= SINH VIÊN TỰ ĐĂNG KÝ HOẠT ĐỘNG =================

@app.route('/tu-dang-ky-hoat-dong/<ma_hd>')
def tu_dang_ky_hoat_dong(ma_hd):
    if 'user_id' not in session:
        return f"<script>alert('Vui lòng đăng nhập!'); window.location.href='/';</script>"
    # Làm mới masv từ database
    if not session.get('masv'):
        try:
            conn0 = get_db_connection()
            cur0 = conn0.cursor(dictionary=True)
            cur0.execute("SELECT MaSV FROM nguoi_dung WHERE MaNguoiDung = %s", (session['user_id'],))
            u0 = cur0.fetchone()
            if u0 and u0.get('MaSV'):
                session['masv'] = u0['MaSV']
            cur0.close(); conn0.close()
        except:
            pass
    if not session.get('masv'):
        return f"<script>alert('Bạn cần đăng nhập bằng tài khoản sinh viên để đăng ký!'); window.location.href='/hoat-dong';</script>"
    ma_sv = session['masv']
    try:
        conn = get_db_connection()
        cursor = conn.cursor(dictionary=True)
        cursor.execute("SELECT * FROM sv_tham_gia_hd WHERE MaSV = %s AND MaHD = %s", (ma_sv, ma_hd))
        if cursor.fetchone():
            cursor.close(); conn.close()
            return f"<script>alert('Bạn đã đăng ký hoạt động này rồi!'); window.location.href='/hoat-dong';</script>"
        cursor.execute("SELECT DiemCong, HocKy, NamHoc FROM hoat_dong WHERE MaHD = %s", (ma_hd,))
        hd = cursor.fetchone()
        if not hd:
            return f"<script>alert('Không tìm thấy hoạt động!'); window.location.href='/hoat-dong';</script>"
        cursor.execute("INSERT INTO sv_tham_gia_hd (MaSV, MaHD, VaiTro, TrangThaiThamGia) VALUES (%s, %s, %s, %s)",
                       (ma_sv, ma_hd, 'Thành viên', 'Đang chờ điểm danh'))
        conn.commit(); cursor.close(); conn.close()
        return f"<script>alert('Đăng ký thành công! Vui lòng chờ BCH xác nhận điểm danh để được cộng điểm ĐRL.'); window.location.href='/hoat-dong';</script>"
    except Exception as e:
        return f"<script>alert('Lỗi: {e}'); window.location.href='/hoat-dong';</script>"


# ================= XUẤT FILE EXCEL =================

def style_excel(ws, headers):
    header_font = Font(bold=True, color='FFFFFF', size=12)
    header_fill = PatternFill(start_color='F97316', end_color='F97316', fill_type='solid')
    header_align = Alignment(horizontal='center', vertical='center')
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
    ws.row_dimensions[1].height = 30


@app.route('/export/sinh-vien')
def export_sinh_vien():
    r = require_login()
    if r: return r
    if session.get('role') == 'sinhvien': return redirect('/sinh-vien')
    try:
        conn = get_db_connection(); cursor = conn.cursor(dictionary=True)
        cursor.execute("SELECT MaSV, HoLotSV, TenSV, TenLop, MaLop, NgaySinh, SDT, Email, QueQuan, LaDangVien, TrangThai FROM sinh_vien ORDER BY STT DESC")
        data = cursor.fetchall(); cursor.close(); conn.close()
    except:
        return "Lỗi dữ liệu"
    wb = Workbook(); ws = wb.active; ws.title = "Sinh vien"
    headers = ['Mã SV', 'Họ lót', 'Tên SV', 'Lớp', 'Mã Lớp', 'Ngày sinh', 'SĐT', 'Email', 'Quê quán', 'Đảng viên', 'Trạng thái']
    style_excel(ws, headers)
    for i, row in enumerate(data, 2):
        ws.cell(row=i, column=1, value=row['MaSV'])
        ws.cell(row=i, column=2, value=row['HoLotSV'])
        ws.cell(row=i, column=3, value=row['TenSV'])
        ws.cell(row=i, column=4, value=row['TenLop'])
        ws.cell(row=i, column=5, value=row['MaLop'])
        ws.cell(row=i, column=6, value=str(row['NgaySinh'] or ''))
        ws.cell(row=i, column=7, value=row['SDT'])
        ws.cell(row=i, column=8, value=row['Email'])
        ws.cell(row=i, column=9, value=row['QueQuan'])
        ws.cell(row=i, column=10, value='Có' if row['LaDangVien'] == 1 else 'Không')
        ws.cell(row=i, column=11, value=row['TrangThai'])
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 30)
    output = io.BytesIO(); wb.save(output); output.seek(0)
    return send_file(output, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name='DanhSachSinhVien.xlsx')


@app.route('/export/hoat-dong')
def export_hoat_dong():
    r = require_login()
    if r: return r
    if session.get('role') == 'sinhvien': return redirect('/hoat-dong')
    try:
        conn = get_db_connection(); cursor = conn.cursor(dictionary=True)
        cursor.execute("""
            SELECT hd.*, l.TenLoaiHD FROM hoat_dong hd
            LEFT JOIN loai_hoat_dong l ON hd.MaLoaiHD = l.MaLoaiHD
            ORDER BY hd.NgayBatDau DESC
        """)
        data = cursor.fetchall(); cursor.close(); conn.close()
    except:
        return "Lỗi dữ liệu"
    wb = Workbook(); ws = wb.active; ws.title = "Hoat dong"
    headers = ['Mã HD', 'Tên hoạt động', 'Loại', 'Ngày BĐ', 'Ngày KT', 'Điểm ĐRL', 'Học kỳ', 'Năm học']
    style_excel(ws, headers)
    for i, row in enumerate(data, 2):
        ws.cell(row=i, column=1, value=row['MaHD'])
        ws.cell(row=i, column=2, value=row['TenHD'])
        ws.cell(row=i, column=3, value=row['TenLoaiHD'])
        ws.cell(row=i, column=4, value=str(row['NgayBatDau'] or ''))
        ws.cell(row=i, column=5, value=str(row['NgayKetThuc'] or ''))
        ws.cell(row=i, column=6, value=row['DiemCong'])
        ws.cell(row=i, column=7, value=row['HocKy'])
        ws.cell(row=i, column=8, value=row['NamHoc'])
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 30)
    output = io.BytesIO(); wb.save(output); output.seek(0)
    return send_file(output, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name='DanhSachHoatDong.xlsx')


@app.route('/export/diem-ren-luyen')
def export_diem_ren_luyen():
    r = require_login()
    if r: return r
    if session.get('role') == 'sinhvien': return redirect('/diem-ren-luyen')
    hoc_ky = request.args.get('HocKy', 'HK2')
    nam_hoc = request.args.get('NamHoc', '')
    nam_hoc_join = " AND d.NamHoc = %s" if nam_hoc else ""
    try:
        conn = get_db_connection(); cursor = conn.cursor(dictionary=True)
        sql = f"""
            SELECT s.MaSV, s.HoLotSV, s.TenSV, s.TenLop, COALESCE(d.Diem, 0) as Diem
            FROM sinh_vien s LEFT JOIN diem_ren_luyen d ON s.MaSV = d.MaSV AND d.HocKy = %s{nam_hoc_join}
            ORDER BY s.TenLop, s.TenSV
        """
        params = [hoc_ky]
        if nam_hoc: params.append(nam_hoc)
        cursor.execute(sql, tuple(params))
        data = cursor.fetchall(); cursor.close(); conn.close()
    except:
        return "Lỗi dữ liệu"
    wb = Workbook(); ws = wb.active; ws.title = "Diem ren luyen"
    headers = ['Mã SV', 'Họ lót', 'Tên SV', 'Lớp', 'Điểm', 'Xếp loại']
    style_excel(ws, headers)
    for i, row in enumerate(data, 2):
        ws.cell(row=i, column=1, value=row['MaSV'])
        ws.cell(row=i, column=2, value=row['HoLotSV'])
        ws.cell(row=i, column=3, value=row['TenSV'])
        ws.cell(row=i, column=4, value=row['TenLop'])
        ws.cell(row=i, column=5, value=row['Diem'])
        d = int(row['Diem'])
        xl = 'Xuất sắc' if d >= 90 else 'Tốt' if d >= 80 else 'Khá' if d >= 65 else 'Trung bình' if d >= 50 else 'Yếu/Kém'
        ws.cell(row=i, column=6, value=xl)
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 30)
    output = io.BytesIO(); wb.save(output); output.seek(0)
    return send_file(output, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name=f'DiemRenLuyen_{hoc_ky}_{nam_hoc}.xlsx')


@app.route('/export/bch')
def export_bch():
    r = require_login()
    if r: return r
    if session.get('role') == 'sinhvien': return redirect('/bch')
    try:
        conn = get_db_connection(); cursor = conn.cursor(dictionary=True)
        cursor.execute("""
            SELECT b.*, s.HoLotSV, s.TenSV, s.TenLop
            FROM bch_doan_hoi b LEFT JOIN sinh_vien s ON b.MaSV = s.MaSV
            ORDER BY b.MaBCH
        """)
        data = cursor.fetchall(); cursor.close(); conn.close()
    except:
        return "Lỗi dữ liệu"
    wb = Workbook(); ws = wb.active; ws.title = "BCH"
    headers = ['Mã BCH', 'Mã SV', 'Họ tên SV', 'Lớp', 'Giáo viên', 'Chức vụ', 'Nhiệm kỳ', 'SĐT']
    style_excel(ws, headers)
    for i, row in enumerate(data, 2):
        ws.cell(row=i, column=1, value=row['MaBCH'])
        ws.cell(row=i, column=2, value=row['MaSV'])
        ws.cell(row=i, column=3, value=f"{row['HoLotSV'] or ''} {row['TenSV'] or ''}".strip())
        ws.cell(row=i, column=4, value=row['TenLop'])
        ws.cell(row=i, column=5, value=row['TenGiaoVien'])
        ws.cell(row=i, column=6, value=row['MaChucVu'])
        ws.cell(row=i, column=7, value=row['NhiemKy'])
        ws.cell(row=i, column=8, value=row['SDT'])
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 30)
    output = io.BytesIO(); wb.save(output); output.seek(0)
    return send_file(output, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name='DanhSachBCH.xlsx')

@app.route('/export/khen-thuong')
def export_khen_thuong():
    r = require_login()
    if r: return r
    if session.get('role') == 'sinhvien': return redirect('/khen-thuong')
    try:
        conn = get_db_connection(); cursor = conn.cursor(dictionary=True)
        cursor.execute("""
            SELECT kt.*, s.HoLotSV, s.TenSV, s.TenLop, dm.TenDanhMuc
            FROM sv_khen_thuong kt
            LEFT JOIN sinh_vien s ON kt.MaSV = s.MaSV
            LEFT JOIN danh_muc_khen_thuong dm ON kt.MaDanhMucKT = dm.MaDanhMucKT
            ORDER BY kt.NgayKhenThuong DESC
        """)
        data = cursor.fetchall(); cursor.close(); conn.close()
    except:
        return "Lỗi dữ liệu"
    wb = Workbook(); ws = wb.active; ws.title = "Khen thuong"
    headers = ['Mã KT', 'Mã SV', 'Sinh viên', 'Lớp', 'Danh mục', 'Ngày', 'Lý do', 'Trạng thái']
    style_excel(ws, headers)
    for i, row in enumerate(data, 2):
        ws.cell(row=i, column=1, value=row['MaKT'])
        ws.cell(row=i, column=2, value=row['MaSV'])
        ws.cell(row=i, column=3, value=f"{row['HoLotSV'] or ''} {row['TenSV'] or ''}".strip())
        ws.cell(row=i, column=4, value=row['TenLop'])
        ws.cell(row=i, column=5, value=row['TenDanhMuc'])
        ws.cell(row=i, column=6, value=str(row['NgayKhenThuong'] or ''))
        ws.cell(row=i, column=7, value=row['LyDo'])
        ws.cell(row=i, column=8, value=row['TrangThai'])
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)
    output = io.BytesIO(); wb.save(output); output.seek(0)
    return send_file(output, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name='DanhSachKhenThuong.xlsx')


# ================= IMPORT EXCEL =================

@app.route('/import/sinh-vien', methods=['POST'])
def import_sinh_vien():
    if 'user_id' not in session or session.get('role') not in ['admin', 'bch']:
        return f"<script>alert('Bạn không có quyền!'); window.location.href='/sinh-vien';</script>"
    if 'file' not in request.files:
        return f"<script>alert('Chưa chọn file!'); window.location.href='/sinh-vien';</script>"
    f = request.files['file']
    if f.filename == '':
        return f"<script>alert('File không hợp lệ!'); window.location.href='/sinh-vien';</script>"
    try:
        wb = load_workbook(f); ws = wb.active
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        conn = get_db_connection(); cursor = conn.cursor()
        cols = ['MaSV', 'HoLotSV', 'TenSV', 'TenLop', 'MaLop', 'NgaySinh', 'SDT', 'Email', 'QueQuan', 'LaDangVien', 'TrangThai']
        for row in rows:
            if not row[0]: continue
            vals = {}
            for i, c in enumerate(cols):
                v = row[i] if i < len(row) else None
                if isinstance(v, str): v = v.strip()
                if v == '' or v is None: v = None
                vals[c] = v
            cursor.execute("SELECT COUNT(*) FROM sinh_vien WHERE MaSV=%s", (vals['MaSV'],))
            exists = cursor.fetchone()[0] > 0
            if not exists:
                cursor.execute("""
                    INSERT INTO sinh_vien (MaSV, HoLotSV, TenSV, TenLop, MaLop, NgaySinh, SDT, Email, QueQuan, LaDangVien, TrangThai)
                    VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
                """, (vals['MaSV'], vals['HoLotSV'], vals['TenSV'], vals['TenLop'], vals['MaLop'],
                      vals['NgaySinh'], vals['SDT'], vals['Email'], vals['QueQuan'],
                      1 if vals['LaDangVien'] in [1, '1', True, 'Có'] else 0,
                      vals['TrangThai'] or 'Đang học'))
        conn.commit(); cursor.close(); conn.close()
        return f"<script>alert('Import thành công!'); window.location.href='/sinh-vien';</script>"
    except Exception as e:
        return f"<script>alert('Lỗi: {e}'); window.location.href='/sinh-vien';</script>"

@app.route('/import/bch', methods=['POST'])
def import_bch():
    if 'user_id' not in session or session.get('role') not in ['admin', 'bch']:
        return f"<script>alert('Bạn không có quyền!'); window.location.href='/bch';</script>"
    if 'file' not in request.files:
        return f"<script>alert('Chưa chọn file!'); window.location.href='/bch';</script>"
    f = request.files['file']
    if f.filename == '':
        return f"<script>alert('File không hợp lệ!'); window.location.href='/bch';</script>"
    try:
        wb = load_workbook(f); ws = wb.active
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        conn = get_db_connection(); cursor = conn.cursor()
        for row in rows:
            if not row[0]: continue
            masv, tengv, machucvu, nhiemky, sdt = [row[i] for i in range(5)]
            if isinstance(masv, str): masv = masv.strip()
            if isinstance(tengv, str): tengv = tengv.strip()
            if isinstance(machucvu, str): machucvu = machucvu.strip()
            if isinstance(nhiemky, str): nhiemky = nhiemky.strip()
            if isinstance(sdt, str): sdt = sdt.strip()
            cursor.execute("INSERT INTO bch_doan_hoi (MaSV, TenGiaoVien, MaChucVu, NhiemKy, SDT) VALUES (%s,%s,%s,%s,%s)",
                           (masv, tengv, machucvu, nhiemky, sdt))
        conn.commit(); cursor.close(); conn.close()
        return f"<script>alert('Import BCH thành công!'); window.location.href='/bch';</script>"
    except Exception as e:
        return f"<script>alert('Lỗi: {e}'); window.location.href='/bch';</script>"

@app.route('/import/diem-ren-luyen', methods=['POST'])
def import_diem_ren_luyen():
    if 'user_id' not in session or session.get('role') not in ['admin', 'bch']:
        return f"<script>alert('Bạn không có quyền!'); window.location.href='/diem-ren-luyen';</script>"
    if 'file' not in request.files:
        return f"<script>alert('Chưa chọn file!'); window.location.href='/diem-ren-luyen';</script>"
    f = request.files['file']
    if f.filename == '':
        return f"<script>alert('File không hợp lệ!'); window.location.href='/diem-ren-luyen';</script>"
    hoc_ky = request.form.get('HocKy', session.get('drl_hocky', 'HK2'))
    nam_hoc = request.form.get('NamHoc', session.get('drl_namhoc', ''))
    if not nam_hoc: nam_hoc = '2025-2026'
    try:
        wb = load_workbook(f); ws = wb.active
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        conn = get_db_connection(); cursor = conn.cursor()
        for row in rows:
            masv, diem = row[0], row[1]
            if not masv: continue
            if isinstance(masv, str): masv = masv.strip()
            try: diem = int(float(str(diem).strip())) if diem is not None else 0
            except: diem = 0
            cursor.execute("SELECT COUNT(*) FROM diem_ren_luyen WHERE MaSV=%s AND HocKy=%s AND NamHoc=%s", (masv, hoc_ky, nam_hoc))
            exists = cursor.fetchone()[0] > 0
            if exists:
                cursor.execute("UPDATE diem_ren_luyen SET Diem=%s WHERE MaSV=%s AND HocKy=%s AND NamHoc=%s", (diem, masv, hoc_ky, nam_hoc))
            else:
                cursor.execute("INSERT INTO diem_ren_luyen (MaSV, HocKy, NamHoc, Diem) VALUES (%s,%s,%s,%s)", (masv, hoc_ky, nam_hoc, diem))
        conn.commit(); cursor.close(); conn.close()
        return f"<script>alert('Import điểm thành công!'); window.location.href='/diem-ren-luyen';</script>"
    except Exception as e:
        return f"<script>alert('Lỗi: {e}'); window.location.href='/diem-ren-luyen';</script>"

# ================= QUẢN LÝ SINH VIÊN (HIỂN THỊ & TÌM KIẾM) =================

@app.route('/sinh-vien')
def sinh_vien():
    r = require_login()
    if r: return r
    
    search_query = request.args.get('search_query', '').strip()
    class_filter = request.args.get('class_filter', '').strip()
    
    try:
        conn = get_db_connection()
        cursor = conn.cursor(dictionary=True)
        
        query = "SELECT * FROM sinh_vien WHERE 1=1"
        params = []
        
        # SV chỉ xem thông tin của mình
        if session.get('role') == 'sinhvien' and session.get('masv'):
            query += " AND MaSV = %s"
            params.append(session['masv'])
        
        if search_query:
            query += " AND (TenSV LIKE %s OR HoLotSV LIKE %s OR MaSV LIKE %s OR TenLop LIKE %s)"
            like_term = f"%{search_query}%"
            params.extend([like_term, like_term, like_term, like_term])
            
        if class_filter:
            query += " AND TenLop = %s"
            params.append(class_filter)
            
        query += " ORDER BY STT DESC"
        cursor.execute(query, tuple(params))
        students = cursor.fetchall()
        
        cursor.execute("SELECT DISTINCT TenLop FROM sinh_vien WHERE TenLop IS NOT NULL AND TenLop != '' ORDER BY TenLop")
        classes = cursor.fetchall()
        
        cursor.close()
        conn.close()
    except Exception as e:
        print(f"Lỗi truy vấn CSDL: {e}")
        students = []
        classes = []
        
    return render_template('sinhvien.html', 
                           students=students, 
                           classes=classes, 
                           search_query=search_query, 
                           class_filter=class_filter)

# ================= THÊM, SỬA, XÓA SINH VIÊN (CÓ CẢNH BÁO TRÙNG MÃ) =================

@app.route('/luu-sinh-vien', methods=['POST'])
def luu_sinh_vien():
    if 'user_id' not in session or session.get('role') != 'admin':
        return f"<script>alert('Bạn không có quyền thực hiện thao tác này!'); window.location.href='/sinh-vien';</script>"
        
    action = request.form.get('action') 
    old_masv = request.form.get('old_MaSV') 
    
    masv = request.form.get('MaSV')
    holot = request.form.get('HoLotSV')
    tensv = request.form.get('TenSV')
    tenlop = request.form.get('TenLop')
    malop = request.form.get('MaLop')
    ngaysinh = request.form.get('NgaySinh')
    if not ngaysinh: ngaysinh = None 
    
    sdt = request.form.get('SDT')
    email = request.form.get('Email')
    quequan = request.form.get('QueQuan')
    dangvien = request.form.get('LaDangVien')
    trangthai = request.form.get('TrangThai')
    
    try:
        conn = get_db_connection()
        cursor = conn.cursor()
        
        if action == 'add':
            cursor.execute("SELECT MaSV FROM sinh_vien WHERE MaSV = %s", (masv,))
            if cursor.fetchone():
                return f"<script>alert('LỖI: Mã sinh viên {masv} đã tồn tại trong hệ thống! Không thể thêm.'); window.location.href='/sinh-vien';</script>"
            
            query = """INSERT INTO sinh_vien 
                       (MaSV, HoLotSV, TenSV, TenLop, MaLop, NgaySinh, SDT, Email, QueQuan, LaDangVien, TrangThai) 
                       VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)"""
            cursor.execute(query, (masv, holot, tensv, tenlop, malop, ngaysinh, sdt, email, quequan, dangvien, trangthai))
            
        elif action == 'edit':
            if masv != old_masv:
                cursor.execute("SELECT MaSV FROM sinh_vien WHERE MaSV = %s", (masv,))
                if cursor.fetchone():
                    return f"<script>alert('LỖI: Mã sinh viên mới {masv} bị trùng với một sinh viên khác!'); window.location.href='/sinh-vien';</script>"
                    
            query = """UPDATE sinh_vien 
                       SET MaSV=%s, HoLotSV=%s, TenSV=%s, TenLop=%s, MaLop=%s, NgaySinh=%s, SDT=%s, Email=%s, QueQuan=%s, LaDangVien=%s, TrangThai=%s 
                       WHERE MaSV=%s"""
            cursor.execute(query, (masv, holot, tensv, tenlop, malop, ngaysinh, sdt, email, quequan, dangvien, trangthai, old_masv))
            
        conn.commit()
        cursor.close()
        conn.close()
        
        return f"<script>alert('Lưu thông tin thành công!'); window.location.href='/sinh-vien';</script>"
        
    except Exception as e:
        return f"<script>alert('Lỗi CSDL: {e}'); window.location.href='/sinh-vien';</script>"


@app.route('/xoa-sinh-vien/<ma_sv>')
def xoa_sinh_vien(ma_sv):
    if 'user_id' not in session or session.get('role') != 'admin':
        return f"<script>alert('Bạn không có quyền thực hiện thao tác này!'); window.location.href='/sinh-vien';</script>"
        
    try:
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute("DELETE FROM sinh_vien WHERE MaSV = %s", (ma_sv,))
        conn.commit()
        cursor.close()
        conn.close()
        return f"<script>alert('Đã xóa sinh viên {ma_sv} thành công!'); window.location.href='/sinh-vien';</script>"
    except Exception as e:
        return f"<script>alert('Lỗi xóa dữ liệu: {e}'); window.location.href='/sinh-vien';</script>"
# ================= QUẢN LÝ BCH ĐOÀN - HỘI =================

@app.route('/bch')
def bch():
    r = require_login()
    if r: return r
    
    # 1. Nhận từ khóa tìm kiếm và bộ lọc chức vụ
    search_query = request.args.get('search_query', '').strip()
    chucvu_filter = request.args.get('chucvu_filter', '').strip()
    
    try:
        conn = get_db_connection()
        cursor = conn.cursor(dictionary=True)
        
        # 2. Dùng LEFT JOIN kết hợp với tìm kiếm động và lọc
        query = """
            SELECT b.*, s.HoLotSV, s.TenSV, s.TenLop 
            FROM bch_doan_hoi b
            LEFT JOIN sinh_vien s ON b.MaSV = s.MaSV
            WHERE 1=1
        """
        params = []
        
        if search_query:
            # Tìm theo Tên SV, Họ lót, Mã SV, Tên Giáo viên hoặc Mã Chức vụ
            query += " AND (s.TenSV LIKE %s OR s.HoLotSV LIKE %s OR b.MaSV LIKE %s OR b.TenGiaoVien LIKE %s OR b.MaChucVu LIKE %s)"
            like_term = f"%{search_query}%"
            params.extend([like_term, like_term, like_term, like_term, like_term])
            
        if chucvu_filter:
            # Lọc chính xác theo Chức vụ
            query += " AND b.MaChucVu = %s"
            params.append(chucvu_filter)
            
        # 3. Đổi thành ASC để sắp xếp tăng dần từ 9, 12, 13... y hệt phpMyAdmin
        query += " ORDER BY b.MaBCH ASC"
        
        cursor.execute(query, tuple(params))
        bch_list = cursor.fetchall()
        
        # 4. Tự động lấy danh sách các chức vụ để đưa vào Menu thả xuống
        cursor.execute("SELECT DISTINCT MaChucVu FROM bch_doan_hoi WHERE MaChucVu IS NOT NULL AND MaChucVu != '' ORDER BY MaChucVu ASC")
        chucvu_list = cursor.fetchall()
        
        cursor.close()
        conn.close()
    except Exception as e:
        print(f"Lỗi truy vấn BCH: {e}")
        bch_list = []
        chucvu_list = []
        
    # Truyền dữ liệu sang giao diện
    return render_template('bch.html', 
                           bch_list=bch_list, 
                           search_query=search_query, 
                           chucvu_filter=chucvu_filter, 
                           chucvu_list=chucvu_list)


@app.route('/luu-bch', methods=['POST'])
def luu_bch():
    if 'user_id' not in session or session.get('role') != 'admin':
        return f"<script>alert('Bạn không có quyền thực hiện thao tác này!'); window.location.href='/bch';</script>"
        
    action = request.form.get('action')
    ma_bch = request.form.get('MaBCH') # Dùng để sửa
    
    masv = request.form.get('MaSV')
    if not masv: masv = None # Nếu để trống thì gán là NULL trong SQL
    
    tengiaovien = request.form.get('TenGiaoVien')
    machucvu = request.form.get('MaChucVu')
    nhiemky = request.form.get('NhiemKy')
    sdt = request.form.get('SDT')
    
    try:
        conn = get_db_connection()
        cursor = conn.cursor()
        
        if action == 'add':
            query = "INSERT INTO bch_doan_hoi (MaSV, TenGiaoVien, MaChucVu, NhiemKy, SDT) VALUES (%s, %s, %s, %s, %s)"
            cursor.execute(query, (masv, tengiaovien, machucvu, nhiemky, sdt))
            if masv:
                cursor.execute("UPDATE nguoi_dung SET VaiTro = 'bch' WHERE MaSV = %s AND VaiTro != 'admin'", (masv,))
        elif action == 'edit':
            query = "UPDATE bch_doan_hoi SET MaSV=%s, TenGiaoVien=%s, MaChucVu=%s, NhiemKy=%s, SDT=%s WHERE MaBCH=%s"
            cursor.execute(query, (masv, tengiaovien, machucvu, nhiemky, sdt, ma_bch))
            
        conn.commit()
        cursor.close()
        conn.close()
        return f"<script>alert('Lưu thông tin BCH thành công!'); window.location.href='/bch';</script>"
    except Exception as e:
        return f"<script>alert('Lỗi CSDL: {e}'); window.location.href='/bch';</script>"


@app.route('/xoa-bch/<int:ma_bch>')
def xoa_bch(ma_bch):
    if 'user_id' not in session or session.get('role') != 'admin':
        return f"<script>alert('Bạn không có quyền thực hiện thao tác này!'); window.location.href='/bch';</script>"
    try:
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute("DELETE FROM bch_doan_hoi WHERE MaBCH = %s", (ma_bch,))
        conn.commit()
        cursor.close()
        conn.close()
        return f"<script>alert('Đã xóa thành công!'); window.location.href='/bch';</script>"
    except Exception as e:
        return f"<script>alert('Lỗi xóa dữ liệu: {e}'); window.location.href='/bch';</script>"
# ================= QUẢN LÝ HOẠT ĐỘNG =================

@app.route('/hoat-dong')
def hoat_dong():
    r = require_login()
    if r: return r
    
    search_query = request.args.get('search_query', '').strip()
    loai_filter = request.args.get('loai_filter', '').strip()
    
    # Làm mới masv từ database cho session hiện tại
    if not session.get('masv'):
        try:
            conn2 = get_db_connection()
            cur2 = conn2.cursor(dictionary=True)
            cur2.execute("SELECT MaSV FROM nguoi_dung WHERE MaNguoiDung = %s", (session['user_id'],))
            u2 = cur2.fetchone()
            if u2 and u2.get('MaSV'):
                session['masv'] = u2['MaSV']
            cur2.close(); conn2.close()
        except:
            pass
    
    try:
        conn = get_db_connection()
        cursor = conn.cursor(dictionary=True)
        
        # 1. Lấy danh sách hoạt động 
        query = """
            SELECT hd.*, l.TenLoaiHD, b.TenGiaoVien, s.TenSV, s.HoLotSV, 
                   b.MaSV as BCH_MaSV, b.MaChucVu
            FROM hoat_dong hd
            LEFT JOIN loai_hoat_dong l ON hd.MaLoaiHD = l.MaLoaiHD
            LEFT JOIN bch_doan_hoi b ON hd.MaBCH = b.MaBCH
            LEFT JOIN sinh_vien s ON b.MaSV = s.MaSV
            WHERE 1=1
        """
        params = []
        if search_query:
            query += " AND (hd.TenHD LIKE %s OR hd.MaHD LIKE %s)"
            params.extend([f"%{search_query}%", f"%{search_query}%"])
        if loai_filter:
            query += " AND hd.MaLoaiHD = %s"
            params.append(loai_filter)
        if session.get('role') == 'sinhvien':
            query += " AND hd.NgayKetThuc >= CURDATE()"
            
        query += " ORDER BY hd.NgayBatDau DESC"
        cursor.execute(query, tuple(params))
        hoatdong_list = cursor.fetchall()
        
        # 2. Lấy dữ liệu phụ trợ
        cursor.execute("SELECT * FROM loai_hoat_dong ORDER BY TenLoaiHD")
        loai_hd_list = cursor.fetchall()
        
        cursor.execute("SELECT b.MaBCH, s.HoLotSV, s.TenSV, b.TenGiaoVien, b.MaChucVu FROM bch_doan_hoi b LEFT JOIN sinh_vien s ON b.MaSV = s.MaSV")
        bch_list = cursor.fetchall()
        
        cursor.execute("SELECT MaSV, HoLotSV, TenSV, TenLop FROM sinh_vien ORDER BY TenLop, TenSV")
        all_students = cursor.fetchall()
        
        # 3. Lấy danh sách SV tham gia từng hoạt động để hiển thị Pop-up chi tiết
        cursor.execute("""
            SELECT stg.MaHD, stg.MaSV, stg.VaiTro, stg.TrangThaiThamGia,
                   stg.MinhChung, stg.NgayDiemDanh,
                   sv.HoLotSV, sv.TenSV, sv.TenLop
            FROM sv_tham_gia_hd stg
            LEFT JOIN sinh_vien sv ON stg.MaSV = sv.MaSV
            ORDER BY stg.MaHD
        """)
        participations = cursor.fetchall()
        
        sv_tham_gia_dict = {}
        for p in participations:
            ma_hd = p['MaHD']
            if ma_hd not in sv_tham_gia_dict:
                sv_tham_gia_dict[ma_hd] = []
            sv_tham_gia_dict[ma_hd].append({
                'MaSV': p['MaSV'],
                'HoTen': f"{p['HoLotSV'] or ''} {p['TenSV'] or ''}".strip(),
                'TenLop': p['TenLop'],
                'TrangThaiThamGia': p['TrangThaiThamGia'],
                'VaiTro': p['VaiTro'],
                'MinhChung': p['MinhChung'],
                'NgayDiemDanh': p['NgayDiemDanh']
            })
        
        my_regs = set()
        if session.get('masv'):
            cursor.execute("SELECT MaHD FROM sv_tham_gia_hd WHERE MaSV = %s", (session['masv'],))
            my_regs = set(r['MaHD'] for r in cursor.fetchall())
        
        cursor.close()
        conn.close()
    except Exception as e:
        print(f"Lỗi truy vấn Hoạt động: {e}")
        hoatdong_list, loai_hd_list, bch_list, all_students, sv_tham_gia_dict, my_regs = [], [], [], [], {}, set()

    return render_template('hoatdong.html', 
                           hoatdong_list=hoatdong_list, loai_hd_list=loai_hd_list,
                           bch_list=bch_list, all_students=all_students,
                           sv_tham_gia_dict=sv_tham_gia_dict,
                           search_query=search_query, loai_filter=loai_filter,
                           my_regs=my_regs)

@app.route('/luu-hoat-dong', methods=['POST'])
def luu_hoat_dong():
    if 'user_id' not in session or session.get('role') not in ['admin', 'bch']:
        return f"<script>alert('Bạn không có quyền thực hiện thao tác này!'); window.location.href='/hoat-dong';</script>"
    
    action = request.form.get('action')
    ma_hd = request.form.get('MaHD')
    ma_loai_hd = request.form.get('MaLoaiHD')
    ten_hd = request.form.get('TenHD')
    ngay_bd = request.form.get('NgayBatDau')
    ngay_kt = request.form.get('NgayKetThuc')
    diem_cong = request.form.get('DiemCong')
    hoc_ky = request.form.get('HocKy')
    nam_hoc = request.form.get('NamHoc')
    ma_bch = request.form.get('MaBCH') or None

    try:
        conn = get_db_connection()
        cursor = conn.cursor()
        
        if action == 'add':
            cursor.execute("SELECT MaHD FROM hoat_dong WHERE MaHD = %s", (ma_hd,))
            if cursor.fetchone():
                return f"<script>alert('LỖI: Mã Hoạt động đã tồn tại!'); window.location.href='/hoat-dong';</script>"
            
            query = "INSERT INTO hoat_dong (MaHD, TenHD, MaLoaiHD, NgayBatDau, NgayKetThuc, DiemCong, MaBCH, HocKy, NamHoc) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)"
            cursor.execute(query, (ma_hd, ten_hd, ma_loai_hd, ngay_bd, ngay_kt, diem_cong, ma_bch, hoc_ky, nam_hoc))
            
        elif action == 'edit':
            query = "UPDATE hoat_dong SET TenHD=%s, MaLoaiHD=%s, NgayBatDau=%s, NgayKetThuc=%s, DiemCong=%s, MaBCH=%s, HocKy=%s, NamHoc=%s WHERE MaHD=%s"
            cursor.execute(query, (ten_hd, ma_loai_hd, ngay_bd, ngay_kt, diem_cong, ma_bch, hoc_ky, nam_hoc, ma_hd))
            
        conn.commit()
        cursor.close()
        conn.close()
        return f"<script>alert('Lưu hoạt động thành công!'); window.location.href='/hoat-dong';</script>"
    except Exception as e:
        print(f"DEBUG LỖI: {e}") 
        return f"<script>alert('Lỗi hệ thống: {e}'); window.history.back();</script>"
    

@app.route('/xoa-hoat-dong/<ma_hd>')
def xoa_hoat_dong(ma_hd):
    if 'user_id' not in session or session.get('role') not in ['admin', 'bch']:
        return f"<script>alert('Bạn không có quyền thực hiện thao tác này!'); window.location.href='/hoat-dong';</script>"
    try:
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute("DELETE FROM hoat_dong WHERE MaHD = %s", (ma_hd,))
        conn.commit()
        cursor.close()
        conn.close()
        return f"<script>alert('Đã xóa hoạt động!'); window.location.href='/hoat-dong';</script>"
    except Exception as e:
        return f"<script>alert('Lỗi xóa: {e}'); window.location.href='/hoat-dong';</script>"

@app.route('/diem-danh-hoat-dong', methods=['POST'])
def diem_danh_hoat_dong():
    if 'user_id' not in session or session.get('role') not in ['admin', 'bch']:
        return f"<script>alert('Bạn không có quyền!'); window.location.href='/hoat-dong';</script>"
    ma_hd = request.form.get('MaHD')
    ds_json = request.form.get('ds_diem_danh', '[]')
    import json
    try:
        danh_sach = json.loads(ds_json)
    except:
        return f"<script>alert('Dữ liệu không hợp lệ!'); window.location.href='/hoat-dong';</script>"
    if not danh_sach:
        return f"<script>alert('Không có dữ liệu điểm danh!'); window.location.href='/hoat-dong';</script>"
    try:
        conn = get_db_connection()
        cursor = conn.cursor(dictionary=True)
        cursor.execute("SELECT DiemCong, HocKy, NamHoc FROM hoat_dong WHERE MaHD = %s", (ma_hd,))
        hd_info = cursor.fetchone()
        if not hd_info:
            return f"<script>alert('Không tìm thấy hoạt động!'); window.location.href='/hoat-dong';</script>"
        diem_cong = int(hd_info['DiemCong'] or 0)
        hoc_ky = hd_info['HocKy'] or 'HK1'
        nam_hoc = hd_info['NamHoc'] or '2025-2026'
        dem_cong = 0; dem_tru = 0
        for sv in danh_sach:
            ma_sv = sv['MaSV']; trang_thai_moi = sv['TrangThai']
            if trang_thai_moi not in ('DaThamGia', 'KhongThamGia'):
                continue
            cursor.execute("SELECT * FROM sv_tham_gia_hd WHERE MaSV = %s AND MaHD = %s", (ma_sv, ma_hd))
            existing = cursor.fetchone()
            trang_thai_cu = existing['TrangThaiThamGia'] if existing else None
            da_co_diem = trang_thai_cu in ('Đã tham gia',)
            if trang_thai_moi == 'DaThamGia':
                if existing:
                    if trang_thai_cu != 'Đã tham gia':
                        cursor.execute("UPDATE sv_tham_gia_hd SET TrangThaiThamGia = 'Đã tham gia' WHERE MaSV = %s AND MaHD = %s", (ma_sv, ma_hd))
                else:
                    cursor.execute("INSERT INTO sv_tham_gia_hd (MaSV, MaHD, VaiTro, TrangThaiThamGia) VALUES (%s, %s, %s, %s)", (ma_sv, ma_hd, 'Thành viên', 'Đã tham gia'))
                if not da_co_diem and diem_cong > 0:
                    cursor.execute("SELECT MaDRL, Diem FROM diem_ren_luyen WHERE MaSV = %s AND HocKy = %s AND NamHoc = %s", (ma_sv, hoc_ky, nam_hoc))
                    drl = cursor.fetchone()
                    if drl:
                        cursor.execute("UPDATE diem_ren_luyen SET Diem = %s WHERE MaDRL = %s", (int(drl['Diem']) + diem_cong, drl['MaDRL']))
                    else:
                        cursor.execute("INSERT INTO diem_ren_luyen (MaSV, HocKy, NamHoc, Diem) VALUES (%s, %s, %s, %s)", (ma_sv, hoc_ky, nam_hoc, diem_cong))
                    dem_cong += 1
            elif trang_thai_moi == 'KhongThamGia':
                if existing:
                    cursor.execute("UPDATE sv_tham_gia_hd SET TrangThaiThamGia = 'Không tham gia' WHERE MaSV = %s AND MaHD = %s", (ma_sv, ma_hd))
                else:
                    cursor.execute("INSERT INTO sv_tham_gia_hd (MaSV, MaHD, VaiTro, TrangThaiThamGia) VALUES (%s, %s, %s, %s)", (ma_sv, ma_hd, 'Thành viên', 'Không tham gia'))
                if da_co_diem and diem_cong > 0:
                    cursor.execute("SELECT MaDRL, Diem FROM diem_ren_luyen WHERE MaSV = %s AND HocKy = %s AND NamHoc = %s", (ma_sv, hoc_ky, nam_hoc))
                    drl = cursor.fetchone()
                    if drl:
                        diem_moi = max(0, int(drl['Diem']) - diem_cong)
                        cursor.execute("UPDATE diem_ren_luyen SET Diem = %s WHERE MaDRL = %s", (diem_moi, drl['MaDRL']))
                    dem_tru += 1
        conn.commit(); cursor.close(); conn.close()
        msg = f'Đã điểm danh xong! {dem_cong} SV được cộng +{diem_cong} ĐRL'
        if dem_tru > 0: msg += f', {dem_tru} SV bị trừ -{diem_cong} ĐRL'
        return f"<script>alert('{msg}'); window.location.href='/hoat-dong';</script>"
    except Exception as e:
        return f"<script>alert('Lỗi: {e}'); window.location.href='/hoat-dong';</script>"


@app.route('/generate-qr/<ma_hd>')
def generate_qr(ma_hd):
    r = require_login()
    if r: return r
    if session.get('role') not in ['admin', 'bch']:
        return f"<script>alert('Bạn không có quyền!'); window.location.href='/hoat-dong';</script>"
    import qrcode, io, base64
    from io import BytesIO
    base_url = request.host_url.rstrip('/')
    qr_url = f"{base_url}/diem-danh-qr/{ma_hd}"
    qr = qrcode.QRCode(box_size=8, border=2)
    qr.add_data(qr_url)
    qr.make(fit=True)
    img = qr.make_image(fill_color="black", back_color="white")
    buf = BytesIO()
    img.save(buf, format='PNG')
    buf.seek(0)
    b64 = base64.b64encode(buf.getvalue()).decode()
    return f'<html><body style="display:flex;justify-content:center;align-items:center;min-height:100vh;margin:0;background:#f1f5f9;font-family:sans-serif;"><div style="background:#fff;padding:40px;border-radius:16px;text-align:center;box-shadow:0 4px 20px rgba(0,0,0,0.1);"><h2 style="margin:0 0 6px;color:#1e293b;">📱 Điểm danh QR</h2><p style="color:#64748b;margin:0 0 20px;font-size:14px;">Hoạt động: <strong>{ma_hd}</strong></p><img src="data:image/png;base64,{b64}" style="width:280px;height:280px;border:4px solid #e2e8f0;border-radius:12px;"><p style="margin-top:16px;font-size:12px;color:#94a3b8;word-break:break-all;">{qr_url}</p><a href="/hoat-dong" style="display:inline-block;margin-top:12px;padding:8px 20px;background:#3b82f6;color:#fff;text-decoration:none;border-radius:8px;font-size:14px;">← Về danh sách</a></div></body></html>'


@app.route('/diem-danh-qr/<ma_hd>', methods=['GET', 'POST'])
def diem_danh_qr(ma_hd):
    r = require_login()
    if r: return r
    if session.get('role') != 'sinhvien':
        return f"<script>alert('Chỉ sinh viên mới có thể điểm danh bằng QR!'); window.location.href='/hoat-dong';</script>"
    if not session.get('masv'):
        try:
            conn0 = get_db_connection()
            cur0 = conn0.cursor(dictionary=True)
            cur0.execute("SELECT MaSV FROM nguoi_dung WHERE MaNguoiDung = %s", (session['user_id'],))
            u0 = cur0.fetchone()
            if u0 and u0.get('MaSV'): session['masv'] = u0['MaSV']
            cur0.close(); conn0.close()
        except:
            pass
    ma_sv = session.get('masv', '')
    if not ma_sv:
        return f"<script>alert('Tài khoản của bạn chưa liên kết với MSSV!'); window.location.href='/hoat-dong';</script>"
    try:
        conn = get_db_connection()
        cursor = conn.cursor(dictionary=True)
        cursor.execute("""SELECT hd.*, l.TenLoaiHD FROM hoat_dong hd LEFT JOIN loai_hoat_dong l ON hd.MaLoaiHD = l.MaLoaiHD WHERE hd.MaHD = %s""", (ma_hd,))
        hd = cursor.fetchone()
        if not hd:
            cursor.close(); conn.close()
            return f"<script>alert('Hoạt động không tồn tại!'); window.location.href='/hoat-dong';</script>"
        cursor.execute("SELECT * FROM sv_tham_gia_hd WHERE MaSV = %s AND MaHD = %s", (ma_sv, ma_hd))
        existing = cursor.fetchone()
        if existing and existing['TrangThaiThamGia'] == 'Đã tham gia':
            cursor.close(); conn.close()
            return f"<script>alert('Bạn đã điểm danh hoạt động này rồi!'); window.location.href='/hoat-dong';</script>"

        diem_cong = int(hd['DiemCong'] or 0)
        hoc_ky = hd['HocKy'] or 'HK1'
        nam_hoc = hd['NamHoc'] or '2025-2026'

        if request.method == 'POST':
            import time, base64, re
            minh_chung = None
            b64_data = request.form.get('minh_chung_b64', '')
            if b64_data:
                match = re.match(r'data:image/(\w+);base64,(.+)', b64_data)
                if match:
                    ext = match.group(1)
                    img_data = base64.b64decode(match.group(2))
                    filename = f"evidence_{ma_sv}_{ma_hd}_{int(time.time())}.{ext}"
                    with open(os.path.join(app.config['UPLOAD_FOLDER'], filename), 'wb') as f:
                        f.write(img_data)
                    minh_chung = filename
            else:
                file_mc = request.files.get('minh_chung')
                if file_mc and file_mc.filename:
                    ext = file_mc.filename.rsplit('.', 1)[-1].lower() if '.' in file_mc.filename else 'jpg'
                    filename = f"evidence_{ma_sv}_{ma_hd}_{int(time.time())}.{ext}"
                    file_mc.save(os.path.join(app.config['UPLOAD_FOLDER'], filename))
                    minh_chung = filename

            from datetime import datetime
            now = datetime.now()
            if existing:
                cursor.execute("UPDATE sv_tham_gia_hd SET TrangThaiThamGia = 'Đã tham gia', MinhChung = COALESCE(%s, MinhChung), NgayDiemDanh = %s WHERE MaSV = %s AND MaHD = %s", (minh_chung, now, ma_sv, ma_hd))
            else:
                cursor.execute("INSERT INTO sv_tham_gia_hd (MaSV, MaHD, VaiTro, TrangThaiThamGia, MinhChung, NgayDiemDanh) VALUES (%s, %s, %s, %s, %s, %s)", (ma_sv, ma_hd, 'Thành viên', 'Đã tham gia', minh_chung, now))
            if diem_cong > 0:
                cursor.execute("SELECT MaDRL, Diem FROM diem_ren_luyen WHERE MaSV = %s AND HocKy = %s AND NamHoc = %s", (ma_sv, hoc_ky, nam_hoc))
                drl = cursor.fetchone()
                if drl:
                    cursor.execute("UPDATE diem_ren_luyen SET Diem = %s WHERE MaDRL = %s", (int(drl['Diem']) + diem_cong, drl['MaDRL']))
                else:
                    cursor.execute("INSERT INTO diem_ren_luyen (MaSV, HocKy, NamHoc, Diem) VALUES (%s, %s, %s, %s)", (ma_sv, hoc_ky, nam_hoc, diem_cong))
            conn.commit(); cursor.close(); conn.close()
            return "<script>alert('✅ Điểm danh thành công! Bạn được cộng +" + str(diem_cong) + " ĐRL.'); window.location.href='/hoat-dong';</script>"

        cursor.close(); conn.close()
        # GET: hiển thị form xác nhận + chụp ảnh minh chứng
        return f"""<!DOCTYPE html>
<html lang="vi">
<head><meta charset="UTF-8"><meta name="viewport" content="width=device-width,initial-scale=1.0">
<title>Điểm danh QR</title><link rel="stylesheet" href="/static/style.css">
<style>
body{{display:flex;justify-content:center;align-items:center;min-height:100vh;margin:0;background:linear-gradient(135deg,#667eea,#764ba2);font-family:sans-serif;}}
.qr-page{{background:#fff;border-radius:16px;padding:32px;max-width:420px;width:90%;box-shadow:0 10px 40px rgba(0,0,0,0.2);text-align:center;}}
.qr-page h2{{color:#1e293b;margin:0 0 20px;}}
.info-row{{display:flex;justify-content:space-between;padding:8px 0;border-bottom:1px solid #f1f5f9;font-size:14px;}}
.info-row .lbl{{color:#64748b;}} .info-row .val{{font-weight:600;color:#1e293b;}}
.camera-box{{margin:20px 0;border:3px dashed #cbd5e1;border-radius:12px;overflow:hidden;background:#f8fafc;position:relative;}}
#camera-feed{{width:100%;max-height:240px;object-fit:cover;display:block;}}
#canvas{{display:none;}}
#captured-preview{{max-width:100%;max-height:240px;display:none;margin:0 auto;border-radius:8px;}}
.btn-group{{display:flex;gap:10px;margin-top:16px;}}
.btn-group .btn{{flex:1;}}
.hidden-input{{display:none;}}
.minhchung-label{{display:inline-block;padding:10px 20px;background:#f1f5f9;border-radius:8px;cursor:pointer;font-size:13px;color:#475569;border:2px dashed #cbd5e1;width:100%;box-sizing:border-box;}}
.minhchung-label:hover{{border-color:#94a3b8;}}
#file-name{{font-size:12px;color:#64748b;margin-top:6px;}}
</style>
</head>
<body>
<div class="qr-page">
<div style="font-size:48px;margin-bottom:8px;">📷</div>
<h2>Điểm danh hoạt động</h2>
<div class="info-row"><span class="lbl">Hoạt động</span><span class="val">{hd['TenHD']}</span></div>
<div class="info-row"><span class="lbl">Loại</span><span class="val">{hd['TenLoaiHD'] or ''}</span></div>
<div class="info-row"><span class="lbl">Sinh viên</span><span class="val">{session.get('username','')} ({ma_sv})</span></div>
<div class="info-row"><span class="lbl">Điểm ĐRL</span><span class="val" style="color:var(--success);">+{diem_cong}</span></div>

<form method="POST" enctype="multipart/form-data" id="qrForm" style="margin-top:20px;">
<div class="camera-box" id="cameraBox">
<video id="camera-feed" autoplay playsinline></video>
<canvas id="canvas"></canvas>
<img id="captured-preview" alt="Ảnh đã chụp">
</div>

<div style="display:flex;gap:10px;margin-bottom:16px;">
<button type="button" class="btn btn-info" id="btnChup" onclick="chupAnh()" style="flex:1;">📸 Chụp ảnh</button>
<label class="btn btn-secondary" style="flex:1;text-align:center;cursor:pointer;">📁 Chọn ảnh
<input type="file" name="minh_chung" accept="image/*" capture="environment" class="hidden-input" id="fileInput" onchange="chonFile(this)">
</label>
</div>

<input type="hidden" name="minh_chung_b64" id="minh_chung_b64" value="">
<div id="file-name"></div>
<button type="submit" class="btn btn-success" style="width:100%;margin-top:16px;padding:14px;">✅ Xác nhận điểm danh</button>
</form>
</div>

<script>
var stream = null;
var video = document.getElementById('camera-feed');
navigator.mediaDevices.getUserMedia({{video:{{facingMode:"environment",width:{{ideal:640}}}},audio:false}})
.then(function(s){{stream=s;video.srcObject=s;}})
.catch(function(e){{document.getElementById('cameraBox').innerHTML='<p style="padding:20px;color:#ef4444;">❌ Không thể mở camera: '+e+'</p>';}});

function chupAnh() {{
var canvas = document.getElementById('canvas');
canvas.width = video.videoWidth; canvas.height = video.videoHeight;
var ctx = canvas.getContext('2d');
ctx.drawImage(video, 0, 0);
var dataUrl = canvas.toDataURL('image/jpeg', 0.8);
document.getElementById('captured-preview').src = dataUrl;
document.getElementById('captured-preview').style.display = 'block';
document.getElementById('minh_chung_b64').value = dataUrl;
document.getElementById('file-name').innerText = '✅ Đã chụp ảnh thành công';
if (stream) {{stream.getTracks().forEach(function(t){{t.stop();}}); stream=null; video.style.display='none';}}
}}

function chonFile(input) {{
var name = input.files.length ? input.files[0].name : '';
document.getElementById('file-name').innerText = name ? '📎 '+name : '';
if (name) {{
var reader = new FileReader();
reader.onload = function(e) {{document.getElementById('captured-preview').src=e.target.result;document.getElementById('captured-preview').style.display='block';}};
reader.readAsDataURL(input.files[0]);
if (stream) {{stream.getTracks().forEach(function(t){{t.stop();}}); stream=null; video.style.display='none';}}
}}
}}

document.getElementById('qrForm').addEventListener('submit', function(e) {{
var b64 = document.getElementById('minh_chung_b64').value;
var file = document.getElementById('fileInput').files.length;
if (!b64 && !file) {{
if (!confirm('Bạn chưa chụp/chọn ảnh minh chứng. Vẫn tiếp tục?')) {{e.preventDefault(); return;}}
}}
}});
</script>
</body>
</html>"""
    except Exception as e:
        return f"<script>alert('Lỗi: {e}'); window.location.href='/hoat-dong';</script>"


@app.route('/luu-sv-tham-gia', methods=['POST'])
def luu_sv_tham_gia():
    if 'user_id' not in session or session.get('role') not in ['admin', 'bch']:
        return f"<script>alert('Bạn không có quyền thực hiện thao tác này!'); window.location.href='/hoat-dong';</script>"
    ma_hd = request.form.get('MaHD')
    danh_sach_sv = request.form.getlist('chon_sv') 
    
    if not danh_sach_sv:
        return f"<script>alert('Vui lòng chọn ít nhất 1 sinh viên!'); window.location.href='/hoat-dong';</script>"

    try:
        conn = get_db_connection()
        cursor = conn.cursor(dictionary=True)
        
        # 1. Lấy thông tin hoạt động từ DB để biết cộng vào kỳ nào
        cursor.execute("SELECT DiemCong, HocKy, NamHoc FROM hoat_dong WHERE MaHD = %s", (ma_hd,))
        hd_info = cursor.fetchone()
        if not hd_info:
            return f"<script>alert('Không tìm thấy hoạt động!'); window.location.href='/hoat-dong';</script>"
            
        diem_cong = hd_info['DiemCong'] or 0
        hoc_ky = hd_info['HocKy'] or 'HK1'
        nam_hoc = hd_info['NamHoc'] or '2025-2026'
        
        sv_thanh_cong = 0
        
        # 2. Lưu SV tham gia (trạng thái chờ điểm danh, chưa cộng điểm)
        for ma_sv in danh_sach_sv:
            cursor.execute("SELECT * FROM sv_tham_gia_hd WHERE MaSV = %s AND MaHD = %s", (ma_sv, ma_hd))
            if not cursor.fetchone():
                cursor.execute("INSERT INTO sv_tham_gia_hd (MaSV, MaHD, VaiTro, TrangThaiThamGia) VALUES (%s, %s, %s, %s)", (ma_sv, ma_hd, 'Thành viên', 'Đang chờ điểm danh'))
                sv_thanh_cong += 1

        conn.commit(); cursor.close(); conn.close()
        return f"<script>alert('Đã thêm thành công {sv_thanh_cong} sinh viên vào danh sách! Vui lòng điểm danh để cộng ĐRL.'); window.location.href='/hoat-dong';</script>"
        
    except Exception as e:
        return f"<script>alert('Lỗi CSDL: {e}'); window.location.href='/hoat-dong';</script>"
   # ================= QUẢN LÝ ĐIỂM RÈN LUYỆN =================

@app.route('/diem-ren-luyen')
def diem_ren_luyen():
    r = require_login()
    if r: return r
    
    hoc_ky = request.args.get('HocKy', session.get('drl_hocky', 'HK2'))
    nam_hoc = request.args.get('NamHoc', '')
    if hoc_ky: session['drl_hocky'] = hoc_ky
    xep_loai = request.args.get('XepLoai', '')
    search_query = request.args.get('search_query', '').strip()
    role = session.get('role')
    
    # Làm mới masv từ database cho session hiện tại
    if not session.get('masv'):
        try:
            conn0 = get_db_connection()
            cur0 = conn0.cursor(dictionary=True)
            cur0.execute("SELECT MaSV FROM nguoi_dung WHERE MaNguoiDung = %s", (session['user_id'],))
            u0 = cur0.fetchone()
            if u0 and u0.get('MaSV'):
                session['masv'] = u0['MaSV']
            cur0.close(); conn0.close()
        except:
            pass
    ma_sv = session.get('masv', '')
    
    nam_hoc_join = " AND d.NamHoc = %s" if nam_hoc else ""

    try:
        conn = get_db_connection()
        cursor = conn.cursor(dictionary=True)
        
        # SV: chỉ xem điểm của mình
        if role == 'sinhvien':
            sql = f"""
                SELECT s.MaSV, s.HoLotSV, s.TenSV, s.TenLop, COALESCE(d.Diem, 0) as Diem
                FROM sinh_vien s
                LEFT JOIN diem_ren_luyen d ON s.MaSV = d.MaSV AND d.HocKy = %s{nam_hoc_join}
                WHERE s.MaSV = %s
            """
            params = [hoc_ky]
            if nam_hoc: params.append(nam_hoc)
            params.append(ma_sv)
            cursor.execute(sql, tuple(params))
            danh_sach = cursor.fetchall()
            view_mode = 'sinh_vien'
            diem_dict = {}
            cursor.close(); conn.close()
            return render_template('diemrenluyen.html', 
                                   danh_sach=danh_sach, view_mode=view_mode, diem_dict=diem_dict,
                                   hoc_ky=hoc_ky, nam_hoc=nam_hoc, xep_loai=xep_loai, search_query=search_query)
        
        # CHẾ ĐỘ 1: Nếu CÓ lọc Xếp loại -> Hiển thị thẳng Danh sách Sinh Viên
        if xep_loai:
            condition = ""
            if xep_loai == 'XuatSac': condition = "COALESCE(d.Diem, 0) >= 90"
            elif xep_loai == 'Tot': condition = "COALESCE(d.Diem, 0) >= 80 AND COALESCE(d.Diem, 0) < 90"
            elif xep_loai == 'Kha': condition = "COALESCE(d.Diem, 0) >= 65 AND COALESCE(d.Diem, 0) < 80"
            elif xep_loai == 'TrungBinh': condition = "COALESCE(d.Diem, 0) >= 50 AND COALESCE(d.Diem, 0) < 65"
            elif xep_loai == 'YeuKem': condition = "COALESCE(d.Diem, 0) < 50"
            
            query = f"""
                SELECT s.MaSV, s.HoLotSV, s.TenSV, s.TenLop, COALESCE(d.Diem, 0) as Diem
                FROM sinh_vien s
                LEFT JOIN diem_ren_luyen d ON s.MaSV = d.MaSV AND d.HocKy = %s{nam_hoc_join}
                WHERE {condition}
            """
            params = [hoc_ky]
            if nam_hoc: params.append(nam_hoc)
            if search_query:
                query += " AND (s.TenLop LIKE %s OR s.TenSV LIKE %s OR s.MaSV LIKE %s)"
                params.extend([f"%{search_query}%", f"%{search_query}%", f"%{search_query}%"])
                
            query += " ORDER BY Diem DESC, s.TenLop, s.TenSV"
            cursor.execute(query, tuple(params))
            danh_sach = cursor.fetchall()
            view_mode = 'sinh_vien'
            diem_dict = {}
            
        # CHẾ ĐỘ 2: Không lọc Xếp loại -> Hiển thị danh sách Lớp
        else:
            if search_query:
                query = f"""
                    SELECT s.MaSV, s.HoLotSV, s.TenSV, s.TenLop, COALESCE(d.Diem, 0) as Diem
                    FROM sinh_vien s
                    LEFT JOIN diem_ren_luyen d ON s.MaSV = d.MaSV AND d.HocKy = %s{nam_hoc_join}
                    WHERE CONCAT(s.HoLotSV, ' ', s.TenSV) LIKE %s OR s.MaSV LIKE %s OR s.TenLop LIKE %s
                    ORDER BY s.TenLop, s.TenSV
                """
                params = [hoc_ky]
                if nam_hoc: params.append(nam_hoc)
                params.extend([f"%{search_query}%" for _ in range(3)])
                cursor.execute(query, tuple(params))
                danh_sach = cursor.fetchall()
                view_mode = 'sinh_vien'
                diem_dict = {}
            else:
                query = "SELECT TenLop, COUNT(MaSV) as SiSo FROM sinh_vien WHERE 1=1"
                params = []
                query += " GROUP BY TenLop ORDER BY TenLop"
                cursor.execute(query, tuple(params))
                danh_sach = cursor.fetchall()
                view_mode = 'lop'
                
                cursor.execute(f"""
                    SELECT s.MaSV, s.HoLotSV, s.TenSV, s.TenLop, COALESCE(d.Diem, 0) as Diem
                    FROM sinh_vien s
                    LEFT JOIN diem_ren_luyen d ON s.MaSV = d.MaSV AND d.HocKy = %s{nam_hoc_join}
                    ORDER BY s.TenLop, s.TenSV
                """, tuple([hoc_ky] + ([nam_hoc] if nam_hoc else [])))
                all_diem = cursor.fetchall()
                
                diem_dict = {}
                for row in all_diem:
                    lop = row['TenLop']
                    if lop not in diem_dict: diem_dict[lop] = []
                    diem = int(row['Diem'])
                    if diem >= 90: xl = "Xuất sắc"
                    elif diem >= 80: xl = "Tốt"
                    elif diem >= 65: xl = "Khá"
                    elif diem >= 50: xl = "Trung bình"
                    else: xl = "Yếu/Kém"
                    row['XepLoai'] = xl
                    diem_dict[lop].append(row)

        cursor.close(); conn.close()
    except Exception as e:
        print(f"Lỗi: {e}"); danh_sach, diem_dict, view_mode = [], {}, 'lop'
        
    return render_template('diemrenluyen.html', 
                           danh_sach=danh_sach, view_mode=view_mode, diem_dict=diem_dict,
                           hoc_ky=hoc_ky, nam_hoc=nam_hoc, xep_loai=xep_loai, search_query=search_query)


@app.route('/cap-nhat-diem', methods=['POST'])
def cap_nhat_diem():
    if 'user_id' not in session or session.get('role') not in ['admin', 'bch']:
        return f"<script>alert('Bạn không có quyền thực hiện thao tác này!'); window.location.href='/diem-ren-luyen';</script>"
    ma_sv, hoc_ky, nam_hoc, diem_moi = request.form.get('MaSV'), request.form.get('HocKy'), request.form.get('NamHoc'), request.form.get('Diem')
    
    try:
        conn = get_db_connection(); cursor = conn.cursor()
        cursor.execute("SELECT MaDRL FROM diem_ren_luyen WHERE MaSV = %s AND HocKy = %s AND NamHoc = %s", (ma_sv, hoc_ky, nam_hoc))
        drl = cursor.fetchone()
        
        if drl:
            cursor.execute("UPDATE diem_ren_luyen SET Diem = %s WHERE MaDRL = %s", (diem_moi, drl[0]))
        else:
            cursor.execute("INSERT INTO diem_ren_luyen (MaSV, HocKy, NamHoc, Diem) VALUES (%s, %s, %s, %s)", (ma_sv, hoc_ky, nam_hoc, diem_moi))
            
        conn.commit(); cursor.close(); conn.close()
        return redirect(f"/diem-ren-luyen?HocKy={hoc_ky}&NamHoc={nam_hoc}")
    except Exception as e:
        return f"<script>alert('Lỗi Cập nhật: {e}'); window.location.href='/diem-ren-luyen';</script>"


# ================= QUẢN LÝ KHEN THƯỞNG =================

@app.route('/khen-thuong')
def khen_thuong():
    r = require_login()
    if r: return r
    
    search_query = request.args.get('search_query', '').strip()
    trangthai_filter = request.args.get('trangthai_filter', '')
    
    khenthuong_list = []
    danhmuc_list = []
    
    # KHỐI 1: Lấy danh mục
    try:
        conn = get_db_connection()
        cursor = conn.cursor(dictionary=True)
        cursor.execute("SELECT * FROM danh_muc_khen_thuong ORDER BY TenDanhMuc")
        danhmuc_list = cursor.fetchall()
        cursor.close(); conn.close()
    except Exception as e:
        print(f"Lỗi lấy danh mục KT: {e}")

    # KHỐI 2: Lấy danh sách (ĐÃ SỬA TÊN BẢNG THÀNH sv_khen_thuong)
    try:
        conn = get_db_connection()
        cursor = conn.cursor(dictionary=True)
        
        query = """
            SELECT kt.*, s.HoLotSV, s.TenSV, s.TenLop, dm.TenDanhMuc
            FROM sv_khen_thuong kt
            LEFT JOIN sinh_vien s ON kt.MaSV = s.MaSV
            LEFT JOIN danh_muc_khen_thuong dm ON kt.MaDanhMucKT = dm.MaDanhMucKT
            WHERE 1=1
        """
        params = []
        if search_query:
            query += " AND (s.TenSV LIKE %s OR kt.MaSV LIKE %s OR kt.LyDo LIKE %s)"
            params.extend([f"%{search_query}%", f"%{search_query}%", f"%{search_query}%"])
        if trangthai_filter:
            query += " AND kt.TrangThai = %s"
            params.append(trangthai_filter)
            
        query += " ORDER BY kt.NgayKhenThuong DESC"
        cursor.execute(query, tuple(params))
        khenthuong_list = cursor.fetchall()
        
        cursor.close(); conn.close()
    except Exception as e:
        print(f"Lỗi lấy danh sách KT: {e}")
        
    return render_template('khenthuong.html', khenthuong_list=khenthuong_list, danhmuc_list=danhmuc_list, search_query=search_query, trangthai_filter=trangthai_filter)


@app.route('/dang-ky-khen-thuong', methods=['POST'])
def dang_ky_khen_thuong():
    if 'user_id' not in session: return redirect(url_for('index'))
    
    masv = request.form.get('MaSV')
    madanhmuc = request.form.get('MaDanhMucKT')
    ngaykt = request.form.get('NgayKhenThuong')
    lydo = request.form.get('LyDo')
    file_minh_chung = request.files.get('FileMinhChung')
    filename = None
    
    if file_minh_chung and file_minh_chung.filename != '':
        filename = secure_filename(file_minh_chung.filename)
        file_minh_chung.save(os.path.join(app.config['UPLOAD_FOLDER'], filename))
    
    try:
        conn = get_db_connection()
        cursor = conn.cursor()
        query = "INSERT INTO sv_khen_thuong (MaSV, MaDanhMucKT, NgayKhenThuong, LyDo, FileMinhChung, TrangThai) VALUES (%s, %s, %s, %s, %s, 'Chờ duyệt')"
        cursor.execute(query, (masv, madanhmuc, ngaykt, lydo, filename))
        conn.commit()
        cursor.close()
        conn.close()
        return f"<script>alert('Đã gửi đề xuất khen thưởng. Vui lòng chờ duyệt!'); window.location.href='/khen-thuong';</script>"
    except Exception as e:
        error_msg = str(e).replace("'", "").replace('"', "").replace("\n", " ")
        return f"<script>alert('Lỗi CSDL: {error_msg}'); window.location.href='/khen-thuong';</script>"


@app.route('/xoa-khen-thuong/<int:ma_kt>')
def xoa_khen_thuong(ma_kt):
    if 'user_id' not in session or session.get('role') not in ['admin', 'bch']:
        return f"<script>alert('Không có quyền!'); window.location.href='/khen-thuong';</script>"
    try:
        conn = get_db_connection(); cursor = conn.cursor()
        cursor.execute("DELETE FROM sv_khen_thuong WHERE MaKT = %s", (ma_kt,))
        conn.commit(); cursor.close(); conn.close()
        return redirect('/khen-thuong')
    except Exception as e:
        return f"<script>alert('Lỗi: {e}'); window.location.href='/khen-thuong';</script>"

@app.route('/duyet-khen-thuong/<int:makt>/<trangthai>')
def duyet_khen_thuong(makt, trangthai):
    if 'user_id' not in session or session.get('role') not in ['admin', 'bch']: 
        return "Không có quyền!", 403
        
    try:
        # ĐÃ SỬA TÊN BẢNG THÀNH sv_khen_thuong
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute("UPDATE sv_khen_thuong SET TrangThai = %s WHERE MaKT = %s", (trangthai, makt))
        conn.commit()
        cursor.close()
        conn.close()
        return redirect('/khen-thuong')
    except Exception as e:
        error_msg = str(e).replace("'", "").replace('"', "").replace("\n", " ")
        return f"<script>alert('Lỗi cập nhật: {error_msg}'); window.location.href='/khen-thuong';</script>"
# ================= QUẢN LÝ TRUYỀN THÔNG =================

@app.route('/truyen-thong')
def truyen_thong():
    r = require_login()
    if r: return r
    
    search_query = request.args.get('search_query', '').strip()
    
    try:
        conn = get_db_connection()
        cursor = conn.cursor(dictionary=True)
        query = "SELECT * FROM truyen_thong WHERE 1=1"
        params = []
        if search_query:
            query += " AND (TieuDe LIKE %s OR NoiDung LIKE %s)"
            like = f"%{search_query}%"
            params.extend([like, like])
        query += " ORDER BY NgayDang DESC"
        cursor.execute(query, tuple(params))
        baiviet_list = cursor.fetchall()
        cursor.close(); conn.close()
        print(f"DEBUG truyen_thong: {len(baiviet_list)} records", flush=True)
    except Exception as e:
        print(f"DEBUG truyen_thong ERROR: {e}", flush=True); baiviet_list = []
        
    return render_template('truyenthong.html', bai_viet_list=baiviet_list, search_query=search_query)

@app.route('/luu-bai-viet', methods=['POST'])
def luu_bai_viet():
    if 'user_id' not in session or session.get('role') not in ['admin', 'bch']:
        return f"<script>alert('Bạn không có quyền thực hiện thao tác này!'); window.location.href='/truyen-thong';</script>"
    
    action = request.form.get('action')
    ma_bv = request.form.get('MaBaiViet')
    tieude = request.form.get('TieuDe')
    noidung = request.form.get('NoiDung')
    ngaydang = request.form.get('NgayDang')
    link = request.form.get('link_bai_viet')
    
    try:
        conn = get_db_connection(); cursor = conn.cursor()
        if action == 'edit':
            cursor.execute("UPDATE truyen_thong SET TieuDe=%s, NoiDung=%s, NgayDang=%s, link_bai_viet=%s WHERE MaBaiViet=%s",
                           (tieude, noidung, ngaydang, link, ma_bv))
        else:
            cursor.execute("INSERT INTO truyen_thong (TieuDe, NoiDung, NgayDang, link_bai_viet) VALUES (%s, %s, %s, %s)", 
                           (tieude, noidung, ngaydang, link))
        conn.commit(); cursor.close(); conn.close()
        return "<script>alert('Đã lưu bài viết!'); window.location.href='/truyen-thong';</script>"
    except Exception as e:
        return f"<script>alert('Lỗi: {e}'); window.location.href='/truyen-thong';</script>"

@app.route('/xoa-bai-viet/<int:id>')
def xoa_bai_viet(id):
    if 'user_id' not in session or session.get('role') not in ['admin', 'bch']:
        return f"<script>alert('Bạn không có quyền thực hiện thao tác này!'); window.location.href='/truyen-thong';</script>"
    try:
        conn = get_db_connection(); cursor = conn.cursor()
        cursor.execute("DELETE FROM truyen_thong WHERE MaBaiViet = %s", (id,))
        conn.commit(); cursor.close(); conn.close()
        return redirect('/truyen-thong')
    except Exception as e:
        return f"<script>alert('Lỗi xóa: {e}');</script>"
@app.route('/export/truyen-thong')
def export_truyen_thong():
    r = require_login()
    if r: return r
    if session.get('role') == 'sinhvien': return redirect('/truyen-thong')
    try:
        conn = get_db_connection(); cursor = conn.cursor(dictionary=True)
        cursor.execute("SELECT * FROM truyen_thong ORDER BY NgayDang DESC")
        data = cursor.fetchall(); cursor.close(); conn.close()
    except:
        return "Lỗi dữ liệu"
    wb = Workbook(); ws = wb.active; ws.title = "Truyen thong"
    headers = ['Mã bài', 'Tiêu đề', 'Nội dung', 'Ngày đăng', 'Link']
    style_excel(ws, headers)
    for i, row in enumerate(data, 2):
        ws.cell(row=i, column=1, value=row['MaBaiViet'])
        ws.cell(row=i, column=2, value=row['TieuDe'])
        ws.cell(row=i, column=3, value=row['NoiDung'])
        ws.cell(row=i, column=4, value=str(row['NgayDang'] or ''))
        ws.cell(row=i, column=5, value=row['link_bai_viet'])
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 60)
    output = io.BytesIO(); wb.save(output); output.seek(0)
    return send_file(output, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name='DanhSachTruyenThong.xlsx')

# ================= QUẢN LÝ ẢNH BÌA & GIỚI THIỆU TRANG CHỦ =================

COVER_IMAGE_PATH = 'static/images/anh-gioi-thieu.jpg'

DEFAULT_INTRO = """<p>Đoàn Thanh niên - Hội Sinh viên Khoa Điện tử, Trường Đại học Kỹ thuật Công nghiệp (TNUT) là tổ chức chính trị - xã hội của sinh viên trong khoa, đóng vai trò quan trọng trong việc giáo dục lý tưởng, đạo đức, lối sống và phát huy tinh thần xung kích, sáng tạo của Đoàn viên thanh niên.</p>
<p>Trong nhiều năm qua, Đoàn - Hội FEE đã tổ chức nhiều hoạt động ý nghĩa như: tình nguyện vì cộng đồng, hiến máu nhân đạo, mùa hè xanh, các cuộc thi học thuật, hoạt động văn hóa - văn nghệ - thể thao, góp phần xây dựng môi trường học tập năng động và sáng tạo cho sinh viên.</p>"""

def get_intro():
    try:
        conn = get_db_connection(); cursor = conn.cursor(dictionary=True)
        cursor.execute("SELECT content FROM site_intro WHERE id=1")
        row = cursor.fetchone(); cursor.close(); conn.close()
        return row['content'] if row and row['content'] else DEFAULT_INTRO
    except: return DEFAULT_INTRO

@app.route('/api/intro', methods=['GET'])
def api_get_intro():
    return get_intro()

@app.route('/api/intro', methods=['POST'])
def api_save_intro():
    if 'user_id' not in session or session.get('role') not in ['admin', 'bch']:
        return 'Không có quyền', 403
    content = request.form.get('content', '')
    try:
        conn = get_db_connection(); cursor = conn.cursor()
        cursor.execute("INSERT INTO site_intro (id, content) VALUES (1, %s) ON DUPLICATE KEY UPDATE content=%s", (content, content))
        conn.commit(); cursor.close(); conn.close()
        return 'OK'
    except Exception as e: return f'Lỗi: {e}', 500

@app.route('/lich-su-dang-nhap')
def lich_su_dang_nhap():
    r = require_login()
    if r: return r
    if session.get('role') != 'admin':
        return f"<script>alert('Chỉ Admin mới có quyền truy cập!'); window.location.href='/home';</script>"
    search_query = request.args.get('search', '').strip()
    status_filter = request.args.get('status', '').strip()
    try:
        conn = get_db_connection(); cursor = conn.cursor(dictionary=True)
        query = "SELECT * FROM login_log WHERE 1=1"
        params = []
        if search_query:
            query += " AND (username LIKE %s OR ip_address LIKE %s)"
            params.extend([f"%{search_query}%", f"%{search_query}%"])
        if status_filter:
            query += " AND trang_thai = %s"
            params.append(status_filter)
        query += " ORDER BY created_at DESC LIMIT 500"
        cursor.execute(query, tuple(params))
        logs = cursor.fetchall(); cursor.close(); conn.close()
    except:
        logs = []
    return render_template('login_history.html', logs=logs, search_query=search_query, status_filter=status_filter)

@app.route('/upload-cover', methods=['POST'])
def upload_cover():
    if 'user_id' not in session or session.get('role') not in ['admin', 'bch']:
        return f"<script>alert('Không có quyền!'); window.location.href='/home';</script>"
    file = request.files.get('cover_image')
    try:
        if file and file.filename:
            file.save(COVER_IMAGE_PATH)
    except Exception as e:
        return f"<script>alert('Lỗi: {e}'); window.location.href='/home';</script>"
    return redirect(url_for('home'))

# ================= KHỞI ĐỘNG SERVER =================
with app.app_context():
    try:
        conn = get_db_connection(); cursor = conn.cursor()
        cursor.execute("CREATE TABLE IF NOT EXISTS site_intro (id INT PRIMARY KEY, content TEXT)")
        cursor.execute("""CREATE TABLE IF NOT EXISTS login_log (
            id INT AUTO_INCREMENT PRIMARY KEY,
            username VARCHAR(50) NOT NULL,
            vai_tro VARCHAR(50) DEFAULT '',
            ho_ten VARCHAR(100) DEFAULT '',
            ip_address VARCHAR(45) DEFAULT '',
            user_agent TEXT,
            trang_thai VARCHAR(20) DEFAULT '',
            chi_tiet TEXT,
            created_at DATETIME DEFAULT CURRENT_TIMESTAMP
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci""")
        try:
            cursor.execute("ALTER TABLE nguoi_dung ADD COLUMN last_login DATETIME DEFAULT NULL")
        except: pass
        conn.commit(); cursor.close(); conn.close()
    except: pass

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5000, debug=True)
    